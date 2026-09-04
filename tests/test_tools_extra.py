"""Tests for tool logic that had no coverage:
music metadata parsing, document/presentation/spreadsheet creation,
contacts cross-script matching, and safe path building."""
from __future__ import annotations

import json
import os
import tempfile
import unittest

from aisha.tools.registry import ToolRegistry


class MusicMetadataTests(unittest.TestCase):
    def test_parse_duration(self):
        self.assertEqual(ToolRegistry._parse_duration("3:15"), 195)
        self.assertEqual(ToolRegistry._parse_duration("1:02:30"), 3750)
        self.assertEqual(ToolRegistry._parse_duration("0:45"), 45)
        self.assertEqual(ToolRegistry._parse_duration(""), 0)
        self.assertEqual(ToolRegistry._parse_duration("LIVE"), 0)

    def test_parse_views(self):
        self.assertEqual(ToolRegistry._parse_views("254,309,642 views"), 254309642)
        self.assertEqual(ToolRegistry._parse_views("1,973 views"), 1973)
        self.assertEqual(ToolRegistry._parse_views("No views"), 0)
        self.assertEqual(ToolRegistry._parse_views(""), 0)


class FileCreationTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp()
        os.environ["AISHA_OUTPUT_DIR"] = self._tmp
        self.reg = ToolRegistry()

    def tearDown(self):
        os.environ.pop("AISHA_OUTPUT_DIR", None)

    def _data(self, raw):
        parsed = json.loads(raw)
        self.assertTrue(parsed.get("ok"), f"tool failed: {parsed}")
        return parsed["data"]

    def test_output_dir_honours_env(self):
        self.assertEqual(os.path.normpath(self.reg._output_dir()),
                         os.path.normpath(self._tmp))

    def test_create_document_writes_a_real_docx(self):
        d = self._data(self.reg.execute("create_document", {
            "title": "Test Essay",
            "sections": [{"heading": "Intro", "body": "Hello world.",
                          "bullets": ["one", "two"]}],
            "open_after": False,
        }))
        self.assertEqual(d["type"], "docx")
        self.assertTrue(os.path.isfile(d["path"]))
        # It must be a valid docx (zip) readable by python-docx.
        from docx import Document
        doc = Document(d["path"])
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Test Essay", text)
        self.assertIn("Hello world.", text)

    def test_append_to_extends_existing_docx(self):
        first = self._data(self.reg.execute("create_document", {
            "title": "Doc A", "sections": [{"heading": "H", "body": "body"}],
            "open_after": False}))
        from docx import Document
        n_before = len(Document(first["path"]).paragraphs)
        upd = self._data(self.reg.execute("create_document", {
            "title": "ignored", "append_to": os.path.basename(first["path"]),
            "sections": [{"heading": "References", "body": "Smith 2024."}],
            "open_after": False}))
        self.assertEqual(upd["status"], "updated")
        n_after = len(Document(upd["path"]).paragraphs)
        self.assertGreater(n_after, n_before)

    def test_create_presentation_writes_pptx_with_slides(self):
        d = self._data(self.reg.execute("create_presentation", {
            "title": "Deck", "theme": "violet",
            "slides": [{"title": "S1", "bullets": ["a", "b"]},
                       {"title": "S2", "bullets": ["c"]}],
            "open_after": False,
        }))
        self.assertEqual(d["type"], "pptx")
        self.assertTrue(os.path.isfile(d["path"]))
        from pptx import Presentation
        prs = Presentation(d["path"])
        # title slide + 2 content slides
        self.assertEqual(len(prs.slides._sldIdLst), 3)

    def test_create_spreadsheet_writes_xlsx(self):
        d = self._data(self.reg.execute("create_spreadsheet", {
            "sheets": [{"name": "Budget", "headers": ["Item", "Cost"],
                        "rows": [["Rent", 15000], ["Food", 8000]]}],
            "open_after": False,
        }))
        self.assertEqual(d["type"], "xlsx")
        self.assertTrue(os.path.isfile(d["path"]))

    def test_write_file_refuses_office_formats(self):
        parsed = json.loads(self.reg.execute("write_file", {
            "path": os.path.join(self._tmp, "x.docx"),
            "content": "nope", "mode": "overwrite"}))
        self.assertFalse(parsed.get("ok"))
        self.assertIn("binary document", str(parsed.get("error", "")))

    def test_unique_path_keeps_unicode_and_avoids_clobber(self):
        p1 = self.reg._unique_path("अनुशासन का महत्व", "docx")
        self.assertIn("अनुशासन का महत्व", os.path.basename(p1))
        open(p1, "w").close()
        p2 = self.reg._unique_path("अनुशासन का महत्व", "docx")
        self.assertNotEqual(p1, p2)  # second call must not collide


class ContactsTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp()
        os.environ["AISHA_DATA_DIR"] = self._tmp
        self.reg = ToolRegistry()

    def tearDown(self):
        os.environ.pop("AISHA_DATA_DIR", None)

    def _d(self, raw):
        return json.loads(raw)["data"]

    def test_lookup_missing_does_not_invent_number(self):
        d = self._d(self.reg.execute("contacts", {"action": "lookup", "name": "Ravi"}))
        self.assertFalse(d.get("found"))
        self.assertNotIn("number", d)

    def test_add_then_cross_script_lookup(self):
        self._d(self.reg.execute("contacts", {
            "action": "add", "name": "सोना", "number": "919812345678"}))
        for q in ["सोना", "sona", "Sona"]:
            d = self._d(self.reg.execute("contacts", {"action": "lookup", "name": q}))
            self.assertTrue(d.get("found"), f"should resolve: {q}")
            self.assertEqual(d.get("number"), "919812345678")


if __name__ == "__main__":
    unittest.main()
